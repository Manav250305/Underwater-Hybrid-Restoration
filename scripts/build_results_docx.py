"""
build_results_docx.py — Task 4: assembles the Review II results section as
a .docx, from the analysis tables produced by build_report_data.py and the
calibration reports written by scripts/calibrate_all.py (stages:
calibrate_routing.py, calibrate_red_compensation.py, calibrate_fusion_alpha.py).

Sections:
  0. Evaluation protocol (held-out split, calibration, metrics, statistics).
  1. Overall comparison table (mean +/- std, all methods, all metrics).
  2. Paired statistical comparison (Holm-adjusted Wilcoxon, effect sizes, CIs).
  3. Where the proposed method loses, and why.
  4. Module 5 — water-type routing incl. red-channel compensation.
  5. Ablation table (Module 5 / Module 7 isolated contributions).
  6. Module 7 — calibrated per-image fusion weight.
  7. Finding: no-reference metrics are misaligned with fidelity for
     choosing fusion weights (why the original Module 7 was replaced).
  8. Downstream validation results.
  9. Known issues & limitations.

Run AFTER the calibration scripts and build_report_data.py (default
--split test) have produced their outputs in --data-dir.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # repo root, so `modules` is importable

import argparse
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

METRICS = ["entropy", "uciqe", "uiqm", "psnr", "ssim"]
METRIC_LABELS = {"entropy": "Entropy", "uciqe": "UCIQE", "uiqm": "UIQM", "psnr": "PSNR (dB)", "ssim": "SSIM"}
PLAIN_LABELS = {"entropy": "entropy", "uciqe": "UCIQE", "uiqm": "UIQM", "psnr": "PSNR", "ssim": "SSIM"}
METHOD_LABELS = {
    "udcp": "UDCP (baseline)",
    "conventional_he": "Conventional HE (baseline)",
    "clahe_only": "CLAHE-only (baseline)",
    "hybrid_neither": "Hybrid — neither module",
    "hybrid_fusionopt_only": "Hybrid — Module 7 only",
    "hybrid_routing_only": "Hybrid — Module 5 only",
    "hybrid_full": "Hybrid — full (proposed)",
    "hybrid_nr_search": "Hybrid — original Module 7 (UCIQE+UIQM search)‡",
}
METHOD_ORDER = ["udcp", "conventional_he", "clahe_only", "hybrid_neither", "hybrid_fusionopt_only",
                "hybrid_routing_only", "hybrid_full", "hybrid_nr_search"]
BASELINES = ["udcp", "conventional_he", "clahe_only"]
TUNING_NOTE = ("All parameter selection (Module 5 profiles, red-compensation strength, Module 7 predictor) "
               "used PSNR/SSIM against references on the calibration split only; every number here is on "
               "the held-out test split. ‡ The original Module 7 maximizes UCIQE+UIQM per image, so its "
               "UCIQE/UIQM values are raised by construction.")


def sig_marker(p):
    if pd.isna(p):
        return ""
    if p < 0.001:
        return "***"
    if p < 0.01:
        return "**"
    if p < 0.05:
        return "*"
    return "ns"


def add_heading(doc, text, level=1, bold=True):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.bold = bold
    return h


def style_table(table):
    table.style = "Light Grid Accent 1"
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def add_table(doc, header, rows):
    table = doc.add_table(rows=1, cols=len(header))
    for i, h in enumerate(header):
        table.rows[0].cells[i].text = h
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    style_table(table)
    return table


def add_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    return p


def fmt_metric(m, v):
    return f"{v:.2f}" if m == "psnr" else f"{v:.3f}"


def fmt_profile(prof):
    return (f"cap={prof['wb_gain_cap']:g}, γ×{prof['gamma_strength']:g}, "
            f"clip×{prof['clahe_clip_scale']:g}, ω={prof['dcp_omega']:g}")


def build_protocol(doc, n_test, n_cal, fixed_gap_db, n_features):
    add_heading(doc, "0. Evaluation Protocol", level=2)
    doc.add_paragraph(
        f"Held-out test split. UIEB-890 was split once, stratified by water type, into a "
        f"calibration split ({n_cal} images, 30%) and a test split ({n_test} images, 70%; seed 42). "
        "Every parameter of the method was chosen on the calibration split in three stages: "
        "(1) per-stage profile multipliers, (2) red-channel compensation strength, and (3) the "
        "Module 7 fusion-weight predictor. Every table in this section is computed on the test "
        "images, which played no part in any of those choices."
    )
    doc.add_paragraph(
        "Fair ablation. Each novel module is compared against an equally tuned alternative, not an "
        "untuned one. Routing OFF uses a single global profile tuned with the same grid, objective and "
        "data as the per-type profiles. Module 7 OFF uses a fixed α=0.5, which was within "
        f"{fixed_gap_db:.2f} dB of the best single fixed α in calibration."
    )
    doc.add_paragraph(
        "Light supervision. Calibration selects 5 scalars per water type and fits one ridge regression "
        f"({n_features} coefficients), using the UIEB reference images of the calibration split. No model is "
        "trained in the deep-learning sense and nothing is learned per pixel, but the method is no "
        "longer strictly reference-free."
    )
    doc.add_paragraph(
        "UCIQE. Computed in true CIELAB, with L*, a* and b* scaled by 1/100 and a bounded saturation "
        "term C/√(C²+L²). Values therefore lie on the usual ≈0–1 scale. Absolute values are not "
        "directly comparable with papers that feed OpenCV's 8-bit Lab encoding into UCIQE unchanged "
        "(Section 9, item 1)."
    )
    doc.add_paragraph(
        "Statistics. Paired two-sided Wilcoxon signed-rank tests, with p-values Holm-adjusted across "
        "every test in the table. With several hundred pairs almost any difference is "
        "'significant', so each comparison also reports the mean paired difference with a bootstrap "
        "95% CI (2,000 resamples) and the matched-pairs rank-biserial correlation r (|r| < 0.1 "
        "negligible, < 0.3 small, < 0.5 medium, otherwise large). Conclusions are drawn from effect "
        "sizes, not from p-values alone."
    )
    add_note(doc, TUNING_NOTE)


def build_overall_comparison(doc, summary_df, n_test):
    add_heading(doc, "1. Overall Comparison — All Methods, All Metrics", level=2)
    doc.add_paragraph(f"Mean ± standard deviation over the {n_test} held-out test images.")
    pm = summary_df.pivot(index="method", columns="metric", values="mean")
    ps = summary_df.pivot(index="method", columns="metric", values="std")
    rows = []
    for method in METHOD_ORDER:
        rows.append([METHOD_LABELS[method]] + [
            f"{fmt_metric(m, pm.loc[method, m])} ± {fmt_metric(m, ps.loc[method, m])}" for m in METRICS])
    add_table(doc, ["Method"] + [METRIC_LABELS[m] for m in METRICS], rows)
    add_note(doc, TUNING_NOTE)


def build_significance(doc, sig_df):
    add_heading(doc, "2. Paired Statistical Comparison — Proposed vs. Each Method", level=2)
    doc.add_paragraph(
        "Δ = hybrid (full) − comparator, so a positive Δ means the proposed method scores higher. "
        "Markers give the Holm-adjusted significance (*** p<0.001, ** p<0.01, * p<0.05, ns not significant)."
    )
    rows = []
    for _, r in sig_df.iterrows():
        comp = r["comparison"].replace("hybrid_full vs ", "")
        digits = 2 if r["metric"] == "psnr" else 3
        rows.append([
            METHOD_LABELS.get(comp, comp), METRIC_LABELS[r["metric"]],
            f"{r['mean_diff']:+.{digits}f} [{r['mean_diff_ci95_lo']:+.{digits}f}, {r['mean_diff_ci95_hi']:+.{digits}f}]",
            f"{r['rank_biserial']:+.2f} ({r['effect_size']})",
            f"{r['p_holm']:.1e} {sig_marker(r['p_holm'])}",
        ])
    add_table(doc, ["Comparator", "Metric", "Mean Δ [95% CI]", "Rank-biserial r", "Holm p"], rows)


def _wins_losses(sig_df, comparator):
    sub = sig_df[sig_df["comparison"] == f"hybrid_full vs {comparator}"]
    wins, losses = [], []
    for _, r in sub.iterrows():
        if not r["significant_holm_0.05"] or r["effect_size"] == "negligible":
            continue
        item = f"{PLAIN_LABELS[r['metric']]} Δ={r['mean_diff']:+.3f} (r={r['rank_biserial']:+.2f}, {r['effect_size']})"
        (wins if r["mean_diff"] > 0 else losses).append(item)
    return wins, losses


def build_loss_discussion(doc, summary_df, sig_df):
    cp = sig_df[(sig_df["comparison"] == "hybrid_full vs clahe_only") & (sig_df["metric"] == "psnr")].iloc[0]
    cs = sig_df[(sig_df["comparison"] == "hybrid_full vs clahe_only") & (sig_df["metric"] == "ssim")].iloc[0]
    if cp["effect_size"] == "negligible" or not cp["significant_holm_0.05"]:
        clahe_sentence = (
            f"Against CLAHE-only, the strongest baseline, the hybrid is level on PSNR (Δ {cp['mean_diff']:+.2f} dB, "
            f"{cp['effect_size']}{'' if cp['significant_holm_0.05'] else ', not significant'}) and higher on SSIM "
            f"(Δ {cs['mean_diff']:+.3f}, {cs['effect_size']} effect). CLAHE-only is a gentle local operator that stays "
            "close to moderate references; the hybrid additionally dehazes and re-balances colour, which improves "
            "structural similarity but can shift global intensity, the quantity PSNR penalizes most.")
    else:
        clahe_sentence = (
            f"Against CLAHE-only, the strongest baseline, the hybrid differs by {cp['mean_diff']:+.2f} dB PSNR "
            f"({cp['effect_size']} effect) and {cs['mean_diff']:+.3f} SSIM ({cs['effect_size']} effect). The DCP "
            "stream dehazes more strongly than a moderate reference, which shifts global intensity: PSNR penalizes "
            "that shift directly, while SSIM, which normalizes local luminance and contrast, penalizes it much less.")
    add_heading(doc, "3. Where the Proposed Method Wins and Loses, and Why", level=2)
    pm = summary_df.pivot(index="method", columns="metric", values="mean")

    def g(method, metric):
        return pm.loc[method, metric]

    doc.add_paragraph(
        "Against each baseline, the lists below give only the differences that are significant after "
        "Holm correction and non-negligible in effect size:"
    )
    for comp in BASELINES:
        wins, losses = _wins_losses(sig_df, comp)
        doc.add_paragraph(
            f"{METHOD_LABELS[comp]} — higher: {'; '.join(wins) or 'none'}. Lower: {'; '.join(losses) or 'none'}.",
            style="List Bullet")

    doc.add_paragraph(
        "PSNR/SSIM against the UIEB references. "
        f"On the test split, hybrid PSNR = {g('hybrid_full','psnr'):.2f} dB and SSIM = "
        f"{g('hybrid_full','ssim'):.3f}, against CLAHE-only {g('clahe_only','psnr'):.2f} dB / "
        f"{g('clahe_only','ssim'):.3f}, conventional HE {g('conventional_he','psnr'):.2f} dB / "
        f"{g('conventional_he','ssim'):.3f}, and UDCP {g('udcp','psnr'):.2f} dB / {g('udcp','ssim'):.3f}. "
        "The UIEB references are not ground-truth clear-water images. Li et al. (2019) produced each "
        "one by having volunteers choose the preferred output among 12 existing enhancement methods. "
        "PSNR/SSIM therefore measure closeness to the output volunteers preferred, which is usually a "
        "moderate correction. " + clahe_sentence
    )
    doc.add_paragraph(
        "UCIQE/entropy vs. conventional HE. "
        f"Conventional HE scores UCIQE = {g('conventional_he','uciqe'):.3f} and entropy = "
        f"{g('conventional_he','entropy'):.2f}, against the hybrid's {g('hybrid_full','uciqe'):.3f} and "
        f"{g('hybrid_full','entropy'):.2f}. Per-channel global equalization stretches every channel to "
        "the full range, which directly inflates luminance contrast, chroma spread and histogram "
        "entropy. Those are the quantities UCIQE and entropy reward, and they rise whether or not the "
        "colours are faithful. Section 7 shows that these no-reference scores do not track reference "
        "fidelity."
    )


def build_module5_section(doc, watertype_df, calib, red, red_effect, out_dir, rounds_text="one round"):
    add_heading(doc, "4. Module 5 — Water-Type Routing with Red-Channel Compensation", level=2)
    doc.add_paragraph(
        "Water absorbs red light first, and how much is lost depends on depth and turbidity. Module 5 "
        "classifies each image's water type from cheap colour and haze statistics, then routes the "
        "per-stage parameters. Those parameters now include the strength of an attenuation-aware "
        "red-channel compensation applied before white balance (Ancuti et al., 2018): "
        "I_r' = I_r + s·(mean I_g − mean I_r)·(1 − I_r)·I_g, which rebuilds red from green where green "
        "is strong and leaves already-red pixels alone."
    )
    doc.add_paragraph(
        f"Stage 1 — profile multipliers. On the calibration split, {calib['n_candidate_settings']} settings "
        "of (WB gain cap, gamma multiplier, CLAHE clip multiplier, DCP ω) were scored at fixed α=0.5 by "
        "J = PSNR/sd + SSIM/sd, with red-channel compensation applied at each arm's current strength. "
        "One global profile (the routing-OFF arm) and one profile per water type were selected. A per-type choice is kept only if it beats the global one on that type's "
        "calibration images (one-sided paired Wilcoxon, p<0.05); otherwise the type falls back to the "
        "global choice."
    )
    gp, gvn = calib["global_profile"], calib["global_vs_neutral_all_calibration"]
    doc.add_paragraph(
        f"Selected global profile: {fmt_profile(gp)}. On calibration images it scores "
        f"{gvn['global_tuned']['psnr']:.2f} dB / {gvn['global_tuned']['ssim']:.3f}, against "
        f"{gvn['neutral']['psnr']:.2f} dB / {gvn['neutral']['ssim']:.3f} for the un-routed defaults."
    )
    rows = []
    for wt, info in calib["types"].items():
        cs = info["calibration_scores"]
        rows.append([wt, info["n_calibration"], fmt_profile(info["hand_set_profile"]),
                     f"{cs['hand_set']['psnr']:.2f} / {cs['hand_set']['ssim']:.3f}",
                     fmt_profile(info["best_per_type_profile"]),
                     f"{cs['per_type_best']['psnr']:.2f} / {cs['per_type_best']['ssim']:.3f}",
                     f"{info['p_better_than_global']:.3g}", info["selected"]])
    add_table(doc, ["Water type", "n (cal.)", "Original hand-set profile", "Hand-set PSNR/SSIM",
                    "Best per-type profile", "Per-type PSNR/SSIM", "p vs. global", "Decision"], rows)

    doc.add_paragraph()
    gs = red["global_strength"]
    gm = red["global_arm_means"]
    doc.add_paragraph(
        f"Stage 2 — red-compensation strength s ∈ {{{', '.join(f'{x:g}' for x in red['strengths_searched'])}}}, "
        "searched with the stage-1 profiles fixed and the same gate. "
        + (f"Stages 1 and 2 were alternated until the selection stopped changing ({rounds_text}). "
           if rounds_text != "a single pass" else "Stage 2 was run once, after stage 1. ")
        + "The global strength is "
        f"s={gs:g}, which scores {gm[str(gs)]['psnr']:.2f} dB / {gm[str(gs)]['ssim']:.3f} on calibration "
        f"images, against {gm['0.0']['psnr']:.2f} dB / {gm['0.0']['ssim']:.3f} with no compensation."
    )
    rows = []
    for wt, info in red["types"].items():
        mb = info["means_by_strength"]
        rows.append([wt, info["n_calibration"], f"{info['best_strength']:g}",
                     f"{mb['0.0']['psnr']:.2f} → {mb[str(info['best_strength'])]['psnr']:.2f}",
                     f"{info['p_better_than_global']:.3g}", f"{info['selected_strength']:g}", info["selected"]])
    add_table(doc, ["Water type", "n (cal.)", "Best s", "PSNR: s=0 → best s", "p vs. global s",
                    "Selected s", "Decision"], rows)
    kept = [wt for wt, info in red["types"].items() if info["selected"] == "per-type"]
    best = ", ".join(f"{wt} {info['best_strength']:g}" for wt, info in red["types"].items())
    add_note(doc, f"Best strength per water type on calibration images: {best}. "
                  + (f"Only {', '.join(kept)} beat the global strength significantly." if kept else
                     f"No type beat the global strength s = {red['global_strength']:g} significantly, so every "
                     "image uses it: the compensation itself matters more than routing its strength."))

    if red_effect is not None:
        doc.add_paragraph()
        doc.add_paragraph(
            "Held-out effect of red-channel compensation. Both arms of the ablation table (Section 5) include "
            "compensation, so its own effect is measured here. The comparison is the 'neither' configuration "
            "(global profile, fixed α=0.5) with and without compensation, paired per test image:"
        )
        rows = []
        for _, r in red_effect.iterrows():
            digits = 2 if r["metric"] == "psnr" else 3
            rows.append([METRIC_LABELS[r["metric"]], fmt_metric(r["metric"], r["mean_without"]),
                         fmt_metric(r["metric"], r["mean_with"]),
                         f"{r['mean_diff']:+.{digits}f} [{r['ci95_lo']:+.{digits}f}, {r['ci95_hi']:+.{digits}f}]",
                         f"{r['rank_biserial']:+.2f} ({r['effect_size']})", f"{r['p_holm']:.1e}"])
        add_table(doc, ["Metric", "Without", "With", "Mean Δ [95% CI]", "Rank-biserial r", "Holm p"], rows)

    doc.add_paragraph()
    doc.add_paragraph(
        "Held-out effect of routing. Below, the full pipeline WITH routing is compared against the "
        "identical pipeline with the global profile (Module 7 on in both arms), per test image, "
        "grouped by assigned water type. Gain = ON − OFF; types that fell back to the global profile "
        "show exactly zero."
    )
    rows = []
    for wt in sorted(watertype_df["water_type"].unique()):
        sub = watertype_df[watertype_df["water_type"] == wt].set_index("metric")
        row = [wt, int(sub.loc["psnr", "n"])]
        for m in ["psnr", "ssim", "uciqe", "uiqm"]:
            digits = 2 if m == "psnr" else 3
            row.append(f"{sub.loc[m, 'mean_routing_gain']:+.{digits}f} ± {sub.loc[m, 'std_routing_gain']:.{digits}f}")
        rows.append(row)
    add_table(doc, ["Water type", "n (test)", "PSNR gain", "SSIM gain", "UCIQE gain", "UIQM gain"], rows)

    sub = watertype_df[watertype_df["metric"].isin(["psnr", "ssim"])]
    types = sorted(sub["water_type"].unique())
    fig, axes = plt.subplots(1, 2, figsize=(8, 3.3))
    for ax, m, unit in zip(axes, ["psnr", "ssim"], [" (dB)", ""]):
        s = sub[sub["metric"] == m].set_index("water_type").loc[types]
        ax.bar(types, s["mean_routing_gain"], color="#4C72B0")
        ax.axhline(0, color="black", linewidth=0.8)
        ax.set_title(f"{PLAIN_LABELS[m]} gain from routing{unit}")
        ax.tick_params(axis="x", rotation=25)
    fig.suptitle("Module 5 effect on held-out test images (routed − global profile)", fontsize=10)
    fig.tight_layout()
    chart_path = os.path.join(out_dir, "_watertype_routing_gain.png")
    fig.savefig(chart_path, dpi=150)
    plt.close(fig)
    doc.add_picture(chart_path, width=Inches(6.0))


def build_ablation_section(doc, ablation_table_df, ablation_deltas_df, n_test):
    add_heading(doc, "5. Ablation Table — Isolating Module 5 and Module 7", level=2)
    doc.add_paragraph(
        f"Mean ± std per metric over the {n_test} test images. Row 3 (neither module: global profile, "
        "fixed α=0.5) is the control, and rows 4–6 isolate each module's individual and combined "
        "contribution. Both arms include the global red-compensation strength, so the routing rows "
        "measure per-type routing only, and the compensation's own effect is reported in Section 4. "
        "Row 7 is the original Module 7 (Section 7)."
    )
    rows = []
    for _, row in ablation_table_df.iterrows():
        rows.append([row["row"]] + [f"{fmt_metric(m, row[f'{m}_mean'])} ± {fmt_metric(m, row[f'{m}_std'])}"
                                    for m in METRICS])
    add_table(doc, ["Configuration"] + [METRIC_LABELS[m] for m in METRICS], rows)

    doc.add_paragraph()
    doc.add_paragraph().add_run("Module contribution (subtraction analysis):").bold = True
    rows = []
    for _, row in ablation_deltas_df.iterrows():
        rows.append([METRIC_LABELS[row["metric"]],
                     f"{row['combined_gain (row6-row3)']:+.3f}",
                     f"{row['fusion_optimizer_alone (row4-row3)']:+.3f}",
                     f"{row['routing_alone (row5-row3)']:+.3f}",
                     f"{row['interaction (row6-row4-row5+row3)']:+.3f}"])
    add_table(doc, ["Metric", "Combined (6−3)", "Module 7 alone (4−3)",
                    "Module 5 alone (5−3)", "Interaction (6−4−5+3)"], rows)
    add_note(doc, TUNING_NOTE)


def build_module7_section(doc, main_df, ablation_deltas_df, alpha_cal, n_features, out_dir):
    add_heading(doc, "6. Module 7 — Calibrated Per-Image Fusion Weight", level=2)
    doc.add_paragraph(
        "The fusion stage blends the CLAHE stream (detail and contrast) with the DCP stream (haze "
        "removal) under a global bias α. Module 7 predicts α per image with a ridge regression from "
        f"{n_features} no-reference statistics of the raw image and of both streams: luminance, contrast, "
        "chroma, colour cast, clipping, haze density and colour dominance. It was fitted on the "
        "calibration split to the per-image α that best matches the reference. At inference it needs "
        "no reference and a single feature pass, and the search loop is gone."
    )
    cv, lam = alpha_cal["cv_by_lambda"], alpha_cal["selected_lambda"]
    best = cv[str(lam)]
    f05, fb, orc = alpha_cal["cv_fixed_alpha_0.5"], alpha_cal["cv_best_fixed_alpha"], alpha_cal["oracle"]
    rows = [
        ["Fixed α = 0.5", f"{f05['psnr']:.2f}", f"{f05['ssim']:.3f}"],
        ["Best single fixed α (chosen inside CV)", f"{fb['psnr']:.2f}", f"{fb['ssim']:.3f}"],
        [f"Predicted α (ridge, λ={lam:g})", f"{best['psnr']:.2f}", f"{best['ssim']:.3f}"],
        ["Oracle per-image α (uses reference; upper bound)", f"{orc['psnr']:.2f}", f"{orc['ssim']:.3f}"],
    ]
    doc.add_paragraph("5-fold cross-validation on the calibration split, grouped by image:")
    add_table(doc, ["α selection", "PSNR", "SSIM"], rows)
    share = (best["psnr"] - f05["psnr"]) / max(orc["psnr"] - f05["psnr"], 1e-9)
    doc.add_paragraph(
        f"The predictor recovers about {100*share:.0f}% of the PSNR gap between a fixed α and the "
        f"per-image oracle (correlation between predicted and oracle α: {best['corr_pred_oracle']:.2f}). "
        "The gain is real but modest, and most of the oracle headroom remains."
    )
    a = main_df["fusion_alpha"].dropna()
    d = ablation_deltas_df.set_index("metric")["fusion_optimizer_alone (row4-row3)"]
    doc.add_paragraph(
        f"Held-out test split. Predicted α has mean {a.mean():.2f} and std {a.std():.2f}, with "
        f"{100*float(((a <= 0.02) | (a >= 0.98)).mean()):.0f}% of images at a boundary. Its isolated "
        f"effect (row 4 − row 3) is PSNR {d['psnr']:+.2f} dB, SSIM {d['ssim']:+.3f}, entropy "
        f"{d['entropy']:+.3f}, UCIQE {d['uciqe']:+.3f}, UIQM {d['uiqm']:+.3f}."
    )
    fig, ax = plt.subplots(figsize=(5, 2.6))
    ax.hist(a, bins=np.linspace(0, 1, 21), color="#4C72B0")
    ax.set_xlabel("predicted fusion α")
    ax.set_ylabel("test images")
    ax.set_title("Module 7: distribution of predicted α (test split)", fontsize=10)
    fig.tight_layout()
    path = os.path.join(out_dir, "_fusion_alpha_hist.png")
    fig.savefig(path, dpi=150)
    plt.close(fig)
    doc.add_picture(path, width=Inches(4.5))


def build_nr_finding_section(doc, align_df, selectors_df, summary_df, sig_df, red_effect, out_dir):
    add_heading(doc, "7. Finding — No-Reference Metrics Are Misaligned with Fidelity for Fusion-Weight Selection", level=2)
    doc.add_paragraph(
        "The original Module 7 chose α per image by maximizing UCIQE+UIQM, a common pattern in "
        "underwater enhancement: optimize the image's own no-reference quality score, with no "
        "reference or training needed. On UIEB this does not work, and the reason is measurable."
    )
    doc.add_paragraph(
        "Within-image alignment (calibration split). For each image, the fused output was computed at "
        "α = 0, 0.1, …, 1 in both pipeline arms. The table gives the median Spearman correlation between "
        "each no-reference score and PSNR/SSIM across those α values, and the share of images where it "
        "is negative, meaning the score ranks α in the opposite order to the reference:"
    )
    rows = [[r["nr_metric"].replace("nr_combined", "UCIQE+UIQM (Module 7 objective)"), r["target"].upper(),
             f"{r['median_spearman']:+.2f}", f"{100*r['frac_negative']:.0f}%"] for _, r in align_df.iterrows()]
    add_table(doc, ["No-reference score", "vs.", "Median within-image Spearman", "Images with negative correlation"], rows)
    doc.add_paragraph()
    doc.add_paragraph("Choosing α by maximizing each score (calibration split):")
    rows = [[r["selector"], f"{r['psnr']:.2f}", f"{r['ssim']:.3f}", f"{100*r['frac_alpha_at_0_or_1']:.0f}%"]
            for _, r in selectors_df.iterrows()]
    add_table(doc, ["α selector", "PSNR", "SSIM", "α at 0 or 1"], rows)

    pm = summary_df.pivot(index="method", columns="metric", values="mean")
    nr = sig_df[sig_df["comparison"] == "hybrid_full vs hybrid_nr_search"].set_index("metric")
    doc.add_paragraph()
    doc.add_paragraph(
        "Held-out confirmation (test split, same routed profile). With α chosen by the UCIQE+UIQM "
        f"search (row 7), PSNR = {pm.loc['hybrid_nr_search','psnr']:.2f} dB and SSIM = "
        f"{pm.loc['hybrid_nr_search','ssim']:.3f}. That compares with "
        f"{pm.loc['hybrid_routing_only','psnr']:.2f} dB / {pm.loc['hybrid_routing_only','ssim']:.3f} at a "
        f"fixed α=0.5 (row 5) and {pm.loc['hybrid_full','psnr']:.2f} dB / {pm.loc['hybrid_full','ssim']:.3f} "
        f"with the predicted α (row 6). The proposed method's advantage over the search is "
        f"Δ PSNR = {nr.loc['psnr','mean_diff']:+.2f} dB (r={nr.loc['psnr','rank_biserial']:+.2f}) and "
        f"Δ SSIM = {nr.loc['ssim','mean_diff']:+.3f} (r={nr.loc['ssim','rank_biserial']:+.2f}), while the "
        f"search raises the score it maximizes (UCIQE {pm.loc['hybrid_nr_search','uciqe']:.3f} vs. "
        f"{pm.loc['hybrid_full','uciqe']:.3f})."
    )
    doc.add_paragraph(
        "Interpretation. Across α, UCIQE and UIQM mostly reward the extremes: the DCP-heavy end "
        "(saturated, high chroma spread) or the CLAHE-heavy end (strong local contrast). The "
        "reference-preferred blend is usually in between. Maximizing these scores therefore selects "
        "a boundary α for most images and reduces fidelity. This is consistent with the observation "
        "that HE scores highest on UCIQE while losing on PSNR/SSIM (Section 3)"
        + (f", and with red-channel compensation, which raises held-out PSNR by "
           f"{red_effect.set_index('metric').loc['psnr','mean_diff']:+.2f} dB while lowering UCIQE by "
           f"{red_effect.set_index('metric').loc['uciqe','mean_diff']:+.3f} and UIQM by "
           f"{red_effect.set_index('metric').loc['uiqm','mean_diff']:+.3f} (Section 4)"
           if red_effect is not None else "")
        + ". The practical "
        "conclusion is that no-reference scores of this family should not be used as optimization "
        "objectives for restoration parameters without first checking their alignment against "
        "references."
    )
    sel = selectors_df.set_index("selector")
    fig, ax = plt.subplots(figsize=(6, 2.8))
    ax.barh(sel.index[::-1], sel["psnr"][::-1], color=["#C44E52" if s.startswith("argmax") else "#4C72B0"
                                                        for s in sel.index[::-1]])
    ax.set_xlabel("PSNR (dB), calibration split")
    ax.set_xlim(sel["psnr"].min() - 0.5, sel["psnr"].max() + 0.3)
    ax.set_title("Choosing α by maximizing a no-reference score vs. fixed / oracle", fontsize=10)
    fig.tight_layout()
    path = os.path.join(out_dir, "_nr_selector_psnr.png")
    fig.savefig(path, dpi=150)
    plt.close(fig)
    doc.add_picture(path, width=Inches(5.5))


def build_downstream_section(doc, downstream_stats):
    add_heading(doc, "8. Downstream Task Validation", level=2)
    doc.add_paragraph(
        "ORB keypoint yield and Canny edge density on the restored vs. raw image, computed with the "
        "corrected detector (nfeatures raised from 1000 to 200,000; see Section 9). Higher counts show "
        "that more structure is detectable. They can also rise from amplified noise or texture, so "
        "they are a proxy for usability, not proof of it."
    )
    add_table(doc, ["Metric", "Value (test split)"], downstream_stats)


def build_limitations(doc, items):
    add_heading(doc, "9. Known Issues & Limitations", level=2)
    doc.add_paragraph(
        "Issues found while building the batch evaluation and the ablation study. Each is recorded "
        "here rather than silently patched, because each would otherwise distort the tables above."
    )
    for title, body, status in items:
        p = doc.add_paragraph()
        p.add_run(f"{title} ({status}). ").bold = True
        p.add_run(body)


def derive_facts(d, calib, alpha_cal, split_df, summary_df, watertype_df):
    """Every current-result number quoted in prose, derived from the output files (never typed in)."""
    from modules.fusion_alpha_predictor import load_model

    pm = summary_df.pivot(index="method", columns="metric", values="mean")
    facts = {
        "fixed_gap_db": abs(alpha_cal["cv_fixed_alpha_0.5"]["psnr"] - alpha_cal["cv_best_fixed_alpha"]["psnr"]),
        "n_features": len(load_model()["features"]),
        "search_cost_db": pm.loc["hybrid_routing_only", "psnr"] - pm.loc["hybrid_nr_search", "psnr"],
        "clear_blue_pct": 100 * (split_df["water_type"] == "clear_blue").mean(),
    }

    grid = pd.read_csv(os.path.join(d, "routing_calibration_grid.csv"))
    if "red_comp_strength" in grid.columns and "red_strengths_used" in calib:  # each arm at its own red strength
        grid = grid[np.isclose(grid["red_comp_strength"], calib["red_strengths_used"]["global"])]
    grid["J"] = grid["psnr"] / calib["sd_psnr"] + grid["ssim"] / calib["sd_ssim"]
    oracle = grid.loc[grid.groupby("filename")["J"].idxmax()].set_index("filename")["psnr"]
    glob = grid
    for k, v in calib["global_profile"].items():
        glob = glob[np.isclose(glob[k], v)]
    facts["oracle_gain_db"] = float((oracle - glob.set_index("filename")["psnr"]).mean())

    w = watertype_df.set_index(["water_type", "metric"])["mean_routing_gain"]
    facts["per_type_gains"] = "; ".join(
        f"{wt} {w[(wt, 'psnr')]:+.2f} dB / {w[(wt, 'ssim')]:+.3f} SSIM"
        for wt in sorted(watertype_df["water_type"].unique()))
    facts["fallback_types"] = ", ".join(wt for wt, info in calib["types"].items() if info["selected"] != "per-type")

    names = {"wb_gain_cap": "WB cap", "gamma_strength": "gamma ×", "clahe_clip_scale": "CLAHE clip ×", "dcp_omega": "DCP ω"}
    # Grid bounds come from the calibration grid that was actually searched (the CSV), not the script's constant
    raw_grid = pd.read_csv(os.path.join(d, "routing_calibration_grid.csv"), usecols=list(names))
    edges = [f"{names[k]}{'' if names[k].endswith('×') else ' '}{v:g}"
             for k, v in calib["global_profile"].items() if k in names and v in (raw_grid[k].min(), raw_grid[k].max())]
    facts["grid_edges"] = ", ".join(edges)

    rounds_path = os.path.join(d, "calibration_rounds.json")
    if os.path.exists(rounds_path):
        with open(rounds_path) as f:
            rr = json.load(f)
        n = len(rr["rounds"])
        facts["rounds_text"] = (("converged: a further pass of both stages left the profiles unchanged" if n == 1
                                 else f"converged after {n} rounds") if rr["converged"]
                                else f"stopped after the maximum of {n} rounds without fully converging")
    else:
        facts["rounds_text"] = "a single pass"
    # Results produced by scripts/calibrate_all.py (which writes calibration_rounds.json) include the gamma fix
    facts["post_gamma_fix"] = os.path.exists(rounds_path)
    sel = pd.read_csv(os.path.join(d, "nr_metric_alpha_selectors.csv")).set_index("selector")
    facts["nr_boundary_pct"] = 100 * float(sel.loc["argmax nr_combined", "frac_alpha_at_0_or_1"])
    return facts


def main():
    parser = argparse.ArgumentParser(description="Build the Review II results section as a .docx")
    parser.add_argument("--data-dir", type=str, default="outputs")
    parser.add_argument("--output", type=str, default="outputs/Review_II_Results_Section.docx")
    args = parser.parse_args()

    d = args.data_dir

    def load_json(name):
        with open(os.path.join(d, name)) as f:
            return json.load(f)

    summary_df = pd.read_csv(os.path.join(d, "summary_stats.csv"))
    sig_df = pd.read_csv(os.path.join(d, "significance_tests.csv"))
    watertype_df = pd.read_csv(os.path.join(d, "per_watertype_breakdown.csv"))
    ablation_table_df = pd.read_csv(os.path.join(d, "ablation_summary.csv"))
    ablation_deltas_df = pd.read_csv(os.path.join(d, "ablation_module_deltas.csv"))
    split_df = pd.read_csv(os.path.join(d, "eval_split.csv"))
    align_df = pd.read_csv(os.path.join(d, "nr_metric_alpha_alignment.csv"))
    selectors_df = pd.read_csv(os.path.join(d, "nr_metric_alpha_selectors.csv"))
    calib = load_json("routing_calibration_selected.json")
    red = load_json("red_comp_calibration_selected.json")
    alpha_cal = load_json("fusion_alpha_calibration.json")
    red_effect_path = os.path.join(d, "red_comp_effect.csv")
    red_effect = pd.read_csv(red_effect_path) if os.path.exists(red_effect_path) else None

    test_names = set(split_df.loc[split_df["split"] == "test", "filename"])
    n_test, n_cal = len(test_names), int((split_df["split"] == "calibration").sum())
    main_df = pd.read_csv(os.path.join(d, "batch_metrics.csv"))
    main_df = main_df[main_df["filename"].isin(test_names)]

    kp = main_df["keypoint_gain_pct"].dropna()
    ed = main_df["edge_density_gain_pct"].dropna()
    cm = main_df["cross_matches"].dropna()
    n_exact_zero_kp = int((kp == 0.0).sum())
    downstream_stats = [
        ("Keypoint gain, mean %", f"{kp.mean():.1f}%"),
        ("Keypoint gain, median %", f"{kp.median():.1f}%"),
        ("Images with exactly 0% keypoint gain", f"{n_exact_zero_kp}/{len(kp)} ({100*n_exact_zero_kp/len(kp):.1f}%)"),
        ("Edge density gain, mean %", f"{ed.mean():.1f}%"),
        ("Edge density gain, median %", f"{ed.median():.1f}%"),
        ("Cross-matches (raw vs. restored), mean", f"{cm.mean():.1f}"),
    ]

    facts = derive_facts(d, calib, alpha_cal, split_df, summary_df, watertype_df)

    limitation_items = [
        (
            "UCIQE computed on OpenCV's 8-bit Lab encoding",
            "compute_uciqe passed cv2.cvtColor(uint8, BGR2LAB) output straight into the UCIQE terms. "
            "That encoding stores L*·255/100 and a*+128, b*+128, so every neutral grey pixel had "
            "'chroma' ≈181, and the unscaled 0–255 luminance range let the contrast term dominate. On a "
            "sample frame, contrast alone contributed 60.4 of a UCIQE of 68.5, and a flat grey image "
            "scored 0.34 instead of 0. The fix converts float [0,1] BGR to true CIELAB, scales L*/a*/b* "
            "by 1/100, centres chroma on the neutral axis, and uses the bounded saturation C/√(C²+L²). "
            "That also removes the near-black C/L blow-up found earlier, which had previously been "
            "patched by dropping pixels with L<1. All UCIQE values are now on the usual ≈0–1 scale. The "
            "earlier '0–100 range' statement was wrong.",
            "Found & Fixed",
        ),
        (
            "Original Module 7 objective was mis-scaled, then shown to be misaligned",
            "With the mis-scaled UCIQE, the UCIQE+UIQM objective came out to ≈49 from UCIQE against ≈0.2 "
            "from UIQM, so it was mostly maximizing a luminance-contrast term. That accident pushed α "
            "toward the CLAHE stream, which made the module look helpful (+0.49 dB in the pre-fix "
            f"ablation). With the metric fixed, the search lowers held-out PSNR by {facts['search_cost_db']:.2f} dB "
            "against a fixed α on the same profile (Section 5, row 7 vs. row 5). Section 7 "
            "shows that the cause is the objective itself rather than a bug, so Module 7 was redesigned "
            "as a calibrated predictor (Section 6). The original search is kept as row 7 of the "
            "ablation for comparison.",
            "Found & Redesigned",
        ),
        (
            "Module 5's white-balance gain cap was never applied",
            "TYPE_PROFILES defined a per-type wb_gain_cap, but restore_image called adaptive_white_balance "
            "without it, so every image used the default cap of 4.0. adaptive_white_balance now takes "
            "max_gain, and both the routed and the global profile pass it.",
            "Found & Fixed",
        ),
        (
            "Routing profiles were hand-set and evaluated on the same images",
            f"The original hand-set profiles made routing hurt clear_blue images, {facts['clear_blue_pct']:.0f}% of UIEB (−1.12 dB "
            "PSNR in the pre-fix all-890 run). Profiles, compensation strength and the α predictor are "
            "now all chosen on a stratified 30% calibration split, against equally tuned alternatives, "
            "with a significance gate that falls back to the global choice. Every reported table uses "
            "the untouched 70% test split.",
            "Found & Fixed",
        ),
        (
            "Per-type routing captures little of the available per-image headroom",
            "On the calibration split, giving each image its own best profile (an oracle that uses the "
            f"reference) would add {facts['oracle_gain_db']:+.2f} dB over the single global profile, but per-type routing recovers "
            "only a small part of it. An exploratory ridge regression from global image statistics, not part of "
            "the released scripts, did no better. Held-out routing gain by water type: "
            f"{facts['per_type_gains']}. Types that fell back to the global profile: "
            f"{facts['fallback_types'] or 'none'}.",
            "Limitation",
        ),
        (
            "Adaptive gamma formula used the opposite convention from how it was applied",
            "compute_adaptive_gamma derives γ = log(0.5)/log(mean) for the convention output = input^γ (γ < 1 "
            "brightens), but adaptive_gamma_correction applied output = input^(1/γ). On its own that "
            "darkened dark images, and darker water types were brightened less than brighter ones. The "
            "earlier calibration had partly compensated by choosing the largest gamma multiplier in its grid. "
            "adaptive_gamma_correction now applies output = input^γ"
            + (", and all three calibration stages and the full evaluation were re-run. Every number in this "
               "section is post-fix." if facts["post_gamma_fix"] else
               ". The numbers in this section were produced before the fix; they will change once "
               "scripts/calibrate_all.py and scripts/run_all_evaluations.sh are re-run."),
            "Found & Fixed" if facts["post_gamma_fix"] else "Fixed in code, re-run pending",
        ),
        (
            "Coordinate-wise, grid-limited calibration",
            (f"Stages 1 and 2 are alternated until the selection stops changing ({facts['rounds_text']}): "
             "a coordinate-wise search, not one joint search over all five parameters, with stage 3 fitted afterwards"
             if facts["post_gamma_fix"] else
             "These results come from a single coordinate-wise pass (stage 1, then stage 2, then stage 3) "
             "rather than a joint search; scripts/calibrate_all.py now alternates stages 1 and 2 until they converge")
            + (f", and some selected global stage-1 values lie at the grid boundary ({facts['grid_edges']})"
               if facts["grid_edges"] else "")
            + ". Stages 1–2 are scored "
            "at fixed α=0.5 and so are not tuned jointly with Module 7.",
            "Limitation",
        ),
        (
            "Significance testing without multiple-comparison correction or effect sizes",
            "The earlier table reported 25 raw Wilcoxon p-values. With n=890, differences as small as "
            "+0.039 UIQM came out at p≈1e-18. The table now carries Holm-adjusted p-values over the whole "
            "family, rank-biserial effect sizes and bootstrap CIs.",
            "Found & Fixed",
        ),
        (
            "ORB keypoint cap artificially tying raw vs. restored counts",
            "cv2.ORB_create(nfeatures=1000) was low enough that raw and restored images routinely hit the "
            "same ceiling, so keypoint_gain_pct read 0% for a large fraction of the dataset. The cap was "
            "raised to nfeatures=200000, since restored images reach up to ~144k keypoints. _pct_change "
            "also returned inf when a raw image had zero keypoints. It now returns None, which is "
            "excluded from aggregates.",
            "Found & Fixed",
        ),
        (
            "Original Module 7 search validity",
            "Before the UCIQE fix, some UCIQE+UIQM(α) curves had two local maxima, so the search gained a "
            "coarse 11-point grid before golden-section refinement. With the corrected metric, a re-scan "
            "across all 4 water types (outputs/investigate_fusion_alpha_postfix.log) found every sampled "
            "curve monotonic, so the search was valid but had no interior optimum to find. Section 7 "
            "explains this behaviour. The full-calibration-split α sweep in Section 7 confirms the result "
            "at scale: maximizing UCIQE+UIQM puts α at 0 or 1 for "
            f"{facts['nr_boundary_pct']:.0f}% of images.",
            "Superseded",
        ),
        (
            "Water-type classifier green_dominance threshold scale bug",
            "The green_turbid rule was gated on green_dominance > 8/255.0, but green_dominance is a "
            "difference of raw 0–255 channel means, so the gate amounted to 'greener than zero'. It was "
            "fixed to > 8, after which 65 of 170 green_turbid images were reclassified.",
            "Found & Fixed",
        ),
        (
            "No learned or recent baselines",
            "The comparison covers classical methods only: UDCP, global HE and CLAHE. Learning-based UIEB "
            "methods such as Water-Net (Li et al., 2019) report substantially higher PSNR/SSIM on UIEB. "
            "This work does not claim to compete with them. Its scope is classical restoration with "
            "light calibration.",
            "Limitation",
        ),
    ]

    doc = Document()
    title = doc.add_heading(level=0)
    title.add_run("Review II — Results Section").bold = True
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run(
        "Hybrid Adaptive Color Correction, CLAHE, and Dehazing Framework for Underwater Image "
        "Restoration and Quality Enhancement — VIT, BCSE403L Digital Image Processing, Fall 2026-2027"
    )
    sub_run.italic = True
    doc.add_paragraph(
        f"Evaluated on the UIEB raw-890 / reference-890 dataset: {n_cal} images for calibration, "
        f"{n_test} held-out images for every reported table. Section numbers (0–9) are relative to "
        "this results section; renumber them to fit the existing Review I document structure."
    )

    build_protocol(doc, n_test, n_cal, facts["fixed_gap_db"], facts["n_features"])
    doc.add_page_break()
    build_overall_comparison(doc, summary_df, n_test)
    doc.add_page_break()
    build_significance(doc, sig_df)
    doc.add_page_break()
    build_loss_discussion(doc, summary_df, sig_df)
    doc.add_page_break()
    build_module5_section(doc, watertype_df, calib, red, red_effect, d, facts["rounds_text"])
    doc.add_page_break()
    build_ablation_section(doc, ablation_table_df, ablation_deltas_df, n_test)
    doc.add_page_break()
    build_module7_section(doc, main_df, ablation_deltas_df, alpha_cal, facts["n_features"], d)
    doc.add_page_break()
    build_nr_finding_section(doc, align_df, selectors_df, summary_df, sig_df, red_effect, d)
    doc.add_page_break()
    build_downstream_section(doc, downstream_stats)
    doc.add_page_break()
    build_limitations(doc, limitation_items)

    doc.save(args.output)
    print(f"Saved -> {args.output}")


if __name__ == "__main__":
    main()
